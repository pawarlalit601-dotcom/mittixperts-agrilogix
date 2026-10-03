"""
AgriLogix Crop Market Category Database Seeder
Seeds crop market categories from the master Excel workbook into the database.
Can be executed with Python:
    python seeds/seed_crop_market_master.py
"""

import os
import sys
from pathlib import Path
import sqlite3
from openpyxl import load_workbook

# Base directory paths
SEEDS_DIR = Path(__file__).resolve().parent
DATABASE_DIR = SEEDS_DIR.parent
DATA_DIR = DATABASE_DIR / "data"
WORKBOOK_PATH = DATA_DIR / "Crop_Market_Category_Master_40_Crops.xlsx"
DB_PATH = DATABASE_DIR / "agrilogix.db"

EXPECTED_HEADERS = (
    "Crop",
    "Market_Category_ID",
    "Valid_Market_Category",
    "Eligibility_Rule",
    "Recommendation_Use",
)


def read_rows() -> list[dict[str, str]]:
    if not WORKBOOK_PATH.is_file():
        raise FileNotFoundError(f"Crop market master workbook is missing: {WORKBOOK_PATH}")
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    if "Crop_Market_Master" not in workbook.sheetnames:
        raise ValueError("Workbook is missing the Crop_Market_Master sheet")
    sheet = workbook["Crop_Market_Master"]
    headers = tuple(str(value or "").strip() for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"Unexpected crop master headers: {headers}")

    rows: list[dict[str, str]] = []
    seen_ids: set[str] = set()
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in values):
            continue
        crop, category_id, category, eligibility, use = (str(value or "").strip() for value in values[:5])
        if not all((crop, category_id, category, eligibility, use)):
            raise ValueError("Crop market master contains an incomplete row")
        if eligibility.upper() not in {"ALLOW", "DISALLOW"}:
            raise ValueError(f"Invalid eligibility rule for {category_id}: {eligibility}")
        category_key = f"{crop.casefold()}::{category_id}"
        if category_key in seen_ids:
            raise ValueError(f"Duplicate market category ID for {crop}: {category_id}")
        seen_ids.add(category_key)
        rows.append({
            "category_id": category_key,
            "source_category_id": category_id,
            "crop": crop,
            "valid_market_category": category,
            "eligibility_rule": eligibility.upper(),
            "recommendation_use": use,
        })
    if not rows:
        raise ValueError("Crop market master contains no category rows")
    return rows


def seed_sqlite(db_file: Path = DB_PATH) -> int:
    if not db_file.exists():
        print(f"Database file {db_file} does not exist. Please run migrations first.")
        return 0

    rows = read_rows()
    conn = sqlite3.connect(db_file)
    cursor = conn.cursor()

    # Ensure table exists
    cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='crop_market_categories';")
    if not cursor.fetchone():
        print("Table 'crop_market_categories' not found. Please run migrations first.")
        conn.close()
        return 0

    cursor.execute("SELECT category_id FROM crop_market_categories")
    existing_ids = {row[0] for row in cursor.fetchall()}

    from datetime import datetime, timezone
    now = datetime.now(timezone.utc).isoformat()

    for row in rows:
        if row["category_id"] in existing_ids:
            cursor.execute(
                """
                UPDATE crop_market_categories
                SET source_category_id = ?, crop = ?, valid_market_category = ?,
                    eligibility_rule = ?, recommendation_use = ?, updated_at = ?
                WHERE category_id = ?
                """,
                (
                    row["source_category_id"],
                    row["crop"],
                    row["valid_market_category"],
                    row["eligibility_rule"],
                    row["recommendation_use"],
                    now,
                    row["category_id"],
                ),
            )
        else:
            cursor.execute(
                """
                INSERT INTO crop_market_categories (
                    category_id, source_category_id, crop, valid_market_category,
                    eligibility_rule, recommendation_use, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    row["category_id"],
                    row["source_category_id"],
                    row["crop"],
                    row["valid_market_category"],
                    row["eligibility_rule"],
                    row["recommendation_use"],
                    now,
                ),
            )

    conn.commit()
    conn.close()
    return len(rows)


if __name__ == "__main__":
    count = seed_sqlite()
    print(f"Successfully seeded {count} crop-market category rows into {DB_PATH.name}.")
