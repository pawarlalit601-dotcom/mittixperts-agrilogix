from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.database import SessionLocal
from app.models import CropMarketCategory


WORKBOOK_PATH = Path(__file__).resolve().parents[1] / "data" / "Crop_Market_Category_Master_40_Crops.xlsx"
EXPECTED_HEADERS = (
    "Crop",
    "Market_Category_ID",
    "Valid_Market_Category",
    "Eligibility_Rule",
    "Recommendation_Use",
)


def read_rows() -> list[dict[str, str]]:
    if not WORKBOOK_PATH.is_file():
        raise FileNotFoundError(f"Crop market master workbook is missing: {WORKBOOK_PATH}")
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    if "Crop_Market_Master" not in workbook.sheetnames:
        raise ValueError("Workbook is missing the Crop_Market_Master sheet")
    sheet = workbook["Crop_Market_Master"]
    headers = tuple(str(value or "").strip() for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"Unexpected crop master headers: {headers}")

    rows: list[dict[str, str]] = []
    seen_ids: set[str] = set()
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in values):
            continue
        crop, category_id, category, eligibility, use = (str(value or "").strip() for value in values[:5])
        if not all((crop, category_id, category, eligibility, use)):
            raise ValueError("Crop market master contains an incomplete row")
        if eligibility.upper() not in {"ALLOW", "DISALLOW"}:
            raise ValueError(f"Invalid eligibility rule for {category_id}: {eligibility}")
        category_key = f"{crop.casefold()}::{category_id}"
        if category_key in seen_ids:
            raise ValueError(f"Duplicate market category ID for {crop}: {category_id}")
        seen_ids.add(category_key)
        rows.append({
            "category_id": category_key,
            "source_category_id": category_id,
            "crop": crop,
            "valid_market_category": category,
            "eligibility_rule": eligibility.upper(),
            "recommendation_use": use,
        })
    if not rows:
        raise ValueError("Crop market master contains no category rows")
    return rows


def seed() -> int:
    rows = read_rows()
    row_ids = {row["category_id"] for row in rows}
    with SessionLocal() as db:
        existing = {item.category_id: item for item in db.scalars(select(CropMarketCategory)).all()}
        for row in rows:
            category = existing.get(row["category_id"])
            if category is None:
                category = CropMarketCategory(**row)
                db.add(category)
            else:
                for field, value in row.items():
                    setattr(category, field, value)
        for category_id, category in existing.items():
            if category_id not in row_ids:
                category.eligibility_rule = "DISALLOW"
        db.commit()
    return len(rows)


if __name__ == "__main__":
    imported = seed()
    print(f"Imported {imported} crop-market category rows from {WORKBOOK_PATH.name}.")